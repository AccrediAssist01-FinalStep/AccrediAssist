"""One-off helper to build copyright code appendix Word file. Not part of runtime app."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AccrediAssist_Copyright_Code_Draft.docx"

SECTIONS: list[tuple[str, str, list[str]]] = [
    (
        "1. WhatsApp Integration (Baileys)",
        "Core WhatsApp module: session persistence, connection lifecycle, QR pairing, "
        "group filtering, message listening, media handling, and HTTP API controllers. "
        "This code implements AccrediAssist institutional WhatsApp intake.",
        [
            "backend/src/whatsapp/session.service.ts",
            "backend/src/whatsapp/whatsapp.config.ts",
            "backend/src/whatsapp/baileys.loader.ts",
            "backend/src/whatsapp/reconnect.service.ts",
            "backend/src/whatsapp/connection.manager.ts",
            "backend/src/whatsapp/whatsapp.service.ts",
            "backend/src/whatsapp/group.filter.ts",
            "backend/src/whatsapp/group.service.ts",
            "backend/src/whatsapp/message.listener.ts",
            "backend/src/whatsapp/message.mapper.ts",
            "backend/src/whatsapp/media.service.ts",
            "backend/src/whatsapp/types.ts",
            "backend/src/controllers/whatsapp.controller.ts",
            "backend/src/routes/whatsapp.routes.ts",
        ],
    ),
    (
        "2. AI Automation Pipeline (Gemini)",
        "Multi-stage AI automation: extraction, classification, validation, duplicate detection, "
        "PDF document processing, and orchestration into pending-review records from WhatsApp messages.",
        [
            "backend/src/ai/services/ai.service.ts",
            "backend/src/ai/providers/gemini.provider.ts",
            "backend/src/ai/services/ai-pipeline.service.ts",
            "backend/src/ai/services/ai-event-report-pipeline.service.ts",
            "backend/src/ai/agents/extraction.agent.ts",
            "backend/src/ai/agents/classification.agent.ts",
            "backend/src/ai/agents/validation.agent.ts",
            "backend/src/ai/agents/duplicate-detection.agent.ts",
            "backend/src/ai/agents/pdf-document.agent.ts",
            "backend/src/ai/agents/ai-event-report.agent.ts",
            "backend/src/ai/utils/ai-config.util.ts",
            "backend/src/ai/utils/prompt-template-config.util.ts",
        ],
    ),
    (
        "3. Backend Server & Workflow Integration",
        "Express application bootstrap, environment validation, health check, server startup with "
        "WhatsApp auto-connect, and workflow service wiring WhatsApp messages into the AI pipeline.",
        [
            "backend/src/server.ts",
            "backend/src/app.ts",
            "backend/src/config/env.ts",
            "backend/src/services/pendingReviewWorkflow.service.ts",
        ],
    ),
    (
        "4. Frontend — WhatsApp Settings UI (reference)",
        "Admin Settings UI for Connect WhatsApp / Scan QR; consumes backend WhatsApp API only "
        "(no session files on the client).",
        [
            "frontend/src/features/settings/components/SettingsWhatsApp.tsx",
            "frontend/src/features/settings/hooks/use-whatsapp-connection.ts",
            "frontend/src/services/whatsapp.service.ts",
        ],
    ),
]


def add_code_block(doc: Document, relative_path: str, content: str) -> None:
    doc.add_heading(relative_path.replace("/", "\\"), level=3)
    p = doc.add_paragraph()
    run = p.add_run(content)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    title = doc.add_heading("AccrediAssist — Copyright Source Code Appendix", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro.add_run(
        "Document purpose: Copyright registration / originality evidence for the AccrediAssist "
        "software (AI-powered academic accreditation management). "
        "The following sections reproduce original source code as maintained in the project repository. "
        "Explanatory notes are provided per section; code is copied verbatim from the codebase."
    )

    meta = doc.add_paragraph()
    meta.add_run(f"Project: AccrediAssist v2.0\n")
    meta.add_run(f"Generated: {date.today().isoformat()}\n")
    meta.add_run("Authors: AccrediAssist development team (Final Step)\n")
    meta.add_run("Repository: AccrediAssist monorepo (backend Express + frontend Next.js)\n")

    doc.add_page_break()

    missing: list[str] = []
    for section_title, explanation, files in SECTIONS:
        doc.add_heading(section_title, level=1)
        doc.add_paragraph(explanation)
        for rel in files:
            path = ROOT / rel.replace("/", "\\")
            if not path.is_file():
                missing.append(rel)
                doc.add_paragraph(f"[MISSING FILE: {rel}]")
                continue
            text = path.read_text(encoding="utf-8")
            add_code_block(doc, rel, text)

    if missing:
        doc.add_heading("Appendix — Missing paths", level=1)
        for m in missing:
            doc.add_paragraph(m)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")
    if missing:
        print("Missing:", ", ".join(missing))


if __name__ == "__main__":
    main()
